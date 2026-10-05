"""Generate a realistic survey .docx (and its raw word/document.xml) so the
Apps Script parser in docx-to-google-form.gs can be checked against genuine
Word XML rather than hand-written XML.

    python tools/make-survey-fixture.py
    node tools/docx-parser-check.js

Writes into build/fixtures/, which is gitignored.
"""
import os
import zipfile

from docx import Document

OUT_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "build", "fixtures"
)
DOCX = os.path.join(OUT_DIR, "survey.docx")
XML = os.path.join(OUT_DIR, "document.xml")


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = Document()

    doc.add_heading("Customer Feedback Survey", level=1)
    doc.add_paragraph(
        "Thanks for taking the time. Your answers help us improve the product."
    )

    doc.add_heading("About you", level=2)

    doc.add_paragraph("1. How satisfied are you with our service?")
    for opt in ["Very satisfied", "Somewhat satisfied", "Neutral", "Dissatisfied"]:
        doc.add_paragraph(opt, style="List Bullet")

    doc.add_paragraph("2. Which of these do you use regularly?")
    for opt in ["Live volume", "Contacts", "Call history"]:
        doc.add_paragraph(opt, style="List Bullet")

    doc.add_paragraph("3. Any comments for the team?")

    doc.add_heading("Preferences", level=2)

    doc.add_paragraph("4. Which integrations do you need? (required)")
    for opt in ["[ ] Slack", "[ ] GitHub", "[x] Jira"]:
        doc.add_paragraph(opt)

    doc.add_paragraph("5. Rate our support on a scale of 1-5")
    doc.add_paragraph("1 - Poor", style="List Bullet")
    doc.add_paragraph("5 - Excellent", style="List Bullet")

    doc.add_paragraph("6. Date of birth")

    doc.add_heading("Final questions", level=2)

    # The trap: numbered options under a numbered question.
    doc.add_paragraph("7. Do you own a car?")
    doc.add_paragraph("1. Yes")
    doc.add_paragraph("2. No")

    doc.add_paragraph("8. Would you recommend us?")

    # An unnumbered, question-mark-only question.
    doc.add_paragraph("Anything else we should know?")

    doc.save(DOCX)

    with zipfile.ZipFile(DOCX) as z:
        data = z.read("word/document.xml")

    with open(XML, "wb") as fh:
        fh.write(data)

    print("wrote " + DOCX)
    print("wrote " + XML + " (%d bytes)" % len(data))


if __name__ == "__main__":
    main()
